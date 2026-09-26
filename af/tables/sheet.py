"""A worksheet as a grid of strings, with the lookups the lab parsers share.

Cells are normalised once here: None -> "", whole floats -> "160", Excel dates -> ISO,
whitespace (including line breaks inside a cell) collapsed. Row and column numbers in
messages are 1-based, as Excel shows them.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

from .rules import RuleTable

_WS = re.compile(r"\s+")


@dataclass
class Sheet:
    path: Path
    name: str
    rows: list[list[str]]

    def cell(self, r: int, c: int) -> str:
        """0-based access; outside the grid is ""."""
        if 0 <= r < len(self.rows) and 0 <= c < len(self.rows[r]):
            return self.rows[r][c]
        return ""

    def where(self, r: int, c: int | None = None) -> str:
        col = "" if c is None else _column_letter(c)
        return f"{self.path.name}[{self.name}]!{col}{r + 1}"

    def find(
        self, pattern: str, *, start: int = 0, stop: int | None = None
    ) -> list[tuple[int, int]]:
        """All (row, col) whose cell fully matches ``pattern`` (case-insensitive)."""
        rx = re.compile(pattern, re.IGNORECASE)
        stop = len(self.rows) if stop is None else stop
        return [
            (r, c)
            for r in range(start, min(stop, len(self.rows)))
            for c, v in enumerate(self.rows[r])
            if v and rx.fullmatch(v)
        ]

    def find_one(self, pattern: str, **kw: int) -> tuple[int, int]:
        hits = self.find(pattern, **kw)
        if len(hits) != 1:
            raise SheetError(
                f"{self.where(0)}: expected one cell matching {pattern!r}, found {len(hits)}"
            )
        return hits[0]


class SheetError(ValueError):
    pass


def apply_cell_fixes(sheet: Sheet, fixes: RuleTable, lab: str) -> list[str]:
    """Apply ``cell_fixes`` rows for this workbook and sheet; return one warning per fix.

    A fix names the cell (``B12``) and what it must hold now (``raw``): when the cell holds
    something else, the workbook has changed under the rule and that is an error, not a
    silent overwrite. Each row is a hand repair someone would otherwise have made in the
    file itself, kept where it can be counted and reviewed."""
    warnings = []
    for rule in fixes.rules:
        if not fixes.in_scope(rule, lab=lab):
            continue
        if rule["file"] != sheet.path.name or rule["sheet"] != sheet.name:
            continue
        r, c = _cell_ref(rule["cell"])
        if sheet.cell(r, c) != rule["raw"]:
            raise SheetError(
                f"{sheet.where(r, c)}: {rule.where} expects {rule['raw']!r}, "
                f"the cell holds {sheet.cell(r, c)!r}"
            )
        while len(sheet.rows) <= r:
            sheet.rows.append([])
        row = sheet.rows[r]
        row.extend([""] * (c + 1 - len(row)))
        row[c] = rule["value"]
        rule.hits += 1
        warnings.append(
            f"{sheet.where(r, c)}: {rule['raw']!r} fixed as {rule['value']!r} by {rule.where}"
        )
    return warnings


def _cell_ref(ref: str) -> tuple[int, int]:
    m = re.fullmatch(r"([A-Z]+)([1-9]\d*)", ref.strip().upper())
    if m is None:
        raise SheetError(f"not a cell reference: {ref!r}")
    c = 0
    for ch in m[1]:
        c = c * 26 + ord(ch) - ord("A") + 1
    return int(m[2]) - 1, c - 1


def load(path: Path) -> list[Sheet]:
    """Every worksheet of the workbook. Needs openpyxl, imported here so that a TSV-only run
    (CDC) does not depend on it."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return [
            Sheet(
                path, ws.title, [[_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            )
            for ws in workbook.worksheets
        ]
    finally:
        workbook.close()


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _WS.sub(" ", str(value)).strip()


def _column_letter(c: int) -> str:
    letters = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters
